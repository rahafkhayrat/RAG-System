from __future__ import annotations

import tempfile
from dataclasses import asdict, dataclass
from pathlib import Path

from ..config import Settings
from ..contracts import Scope
from ..embeddings.base import Embedder
from ..embeddings.hashing import HashingEmbedder
from ..generation.llm import LLMClient, OpenAICompatibleLLM
from ..pipeline import RagEngine
from ..store.memory import InMemoryVectorStore
from .metrics import hit_at_k, recall_at_k, token_f1


@dataclass(frozen=True)
class ToyDoc:
    document_id: str
    course_id: str
    filename: str
    blocks: tuple[tuple[str, str], ...]


@dataclass(frozen=True)
class BenchQuestion:
    qid: str
    course_id: str
    question: str
    gold_docs: tuple[str, ...]
    language: str
    case_type: str
    reference: str = ""
    expect_abstain: bool = False


TOY_DOCS: tuple[ToyDoc, ...] = (
    ToyDoc(
        document_id="algebra_en",
        course_id="math101",
        filename="algebra_en.docx",
        blocks=(
            ("Heading 1", "Linear Equations"),
            (
                "Normal",
                "A linear equation has the general form a x plus b equals zero, where a is not zero. "
                "To solve it, isolate the variable on one side and divide by its coefficient.",
            ),
            ("Heading 1", "Quadratic Equations"),
            (
                "Normal",
                "A quadratic equation has the form a x squared plus b x plus c equals zero. "
                "Its solutions are given by the quadratic formula. The discriminant b squared minus "
                "four a c determines the number of real roots.",
            ),
            ("Heading 1", "Systems of Equations"),
            (
                "Normal",
                "A system of linear equations can be solved by substitution, elimination, or by "
                "using determinants.",
            ),
        ),
    ),
    ToyDoc(
        document_id="algebra_ar",
        course_id="math101",
        filename="algebra_ar.docx",
        blocks=(
            ("Heading 1", "المعادلات الخطية"),
            (
                "Normal",
                "المعادلة الخطية لها الصورة العامة أ س زائد ب يساوي صفر، حيث أ لا يساوي صفراً. "
                "لحلها نعزل المتغير في طرف واحد ثم نقسم على معامله.",
            ),
            ("Heading 1", "المعادلات التربيعية"),
            (
                "Normal",
                "المعادلة التربيعية على الصورة أ س تربيع زائد ب س زائد ج يساوي صفر. "
                "المميز يحدد عدد الجذور الحقيقية.",
            ),
            ("Heading 1", "قاعدة كرامر"),
            (
                "Normal",
                "تنص قاعدة كرامر على إمكانية حل الأنظمة الخطية باستخدام المحددات، "
                "حيث تُستبدل أعمدة المصفوفة بأعمدة الحدود الثابتة.",
            ),
        ),
    ),
    ToyDoc(
        document_id="history_en",
        course_id="hist101",
        filename="history_en.docx",
        blocks=(
            ("Heading 1", "Industrial Revolution"),
            (
                "Normal",
                "The Industrial Revolution began in Britain in the late eighteenth century and "
                "transformed manufacturing through mechanization and steam power.",
            ),
        ),
    ),
)

QUESTIONS: tuple[BenchQuestion, ...] = (
    BenchQuestion(
        qid="q1",
        course_id="math101",
        question="What is the general form of a linear equation?",
        gold_docs=("algebra_en",),
        language="en",
        case_type="full_support",
        reference="a x plus b equals zero",
    ),
    BenchQuestion(
        qid="q2",
        course_id="math101",
        question="ما الصورة العامة للمعادلة الخطية؟",
        gold_docs=("algebra_ar",),
        language="ar",
        case_type="full_support",
        reference="أ س زائد ب يساوي صفر",
    ),
    BenchQuestion(
        qid="q3",
        course_id="math101",
        question="What does the discriminant determine?",
        gold_docs=("algebra_en",),
        language="en",
        case_type="full_support",
        reference="the number of real roots",
    ),
    BenchQuestion(
        qid="q4",
        course_id="math101",
        question="What is Cramer's rule used for?",
        gold_docs=("algebra_ar",),
        language="en",
        case_type="cross_lingual",
        reference="solving linear systems using determinants",
    ),
    BenchQuestion(
        qid="q5",
        course_id="math101",
        question="على ماذا تعتمد قاعدة كرامر؟",
        gold_docs=("algebra_ar",),
        language="ar",
        case_type="full_support",
        reference="المحددات",
    ),
    BenchQuestion(
        qid="q6",
        course_id="hist101",
        question="When did the Industrial Revolution begin?",
        gold_docs=("history_en",),
        language="en",
        case_type="full_support",
        reference="late eighteenth century in Britain",
    ),
    BenchQuestion(
        qid="q7",
        course_id="math101",
        question="When did the Industrial Revolution begin?",
        gold_docs=(),
        language="en",
        case_type="missing_evidence",
        expect_abstain=True,
    ),
    BenchQuestion(
        qid="q8",
        course_id="math101",
        question="Why is the equation x squared plus x equals zero a linear equation?",
        gold_docs=(),
        language="en",
        case_type="false_premise",
        expect_abstain=True,
    ),
)


def _write_docx(path: Path, blocks: tuple[tuple[str, str], ...]) -> None:
    import docx

    document = docx.Document()
    for style, text in blocks:
        if style.startswith("Heading"):
            document.add_heading(text, level=int(style.split()[-1]))
        else:
            document.add_paragraph(text)
    document.save(str(path))


def _make_embedder(name: str) -> tuple[Embedder, str]:
    if name == "bge-m3":
        from ..embeddings.bge_m3 import BgeM3Embedder

        return BgeM3Embedder(), "bge-m3"
    return HashingEmbedder(dim=256), "hashing"


def _make_llm(use_llm: bool) -> tuple[LLMClient | None, str]:
    if not use_llm:
        return None, "disabled"
    settings = Settings.from_env()
    if not settings.has_llm():
        return None, "no-key"
    llm = OpenAICompatibleLLM(settings.llm_base_url, settings.llm_api_key, settings.llm_model)
    return llm, settings.llm_model


def run_benchmark(
    *,
    embedder_name: str = "hashing",
    top_k: int = 3,
    use_llm: bool = False,
    tenant_id: str = "bench_tenant",
) -> dict:
    embedder, embedder_label = _make_embedder(embedder_name)
    llm, llm_label = _make_llm(use_llm)
    engine = RagEngine(
        store=InMemoryVectorStore(),
        embedder=embedder,
        llm=llm,
        top_k=top_k,
    )

    rows: list[dict] = []
    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)
        for toy in TOY_DOCS:
            path = tmpdir / toy.filename
            _write_docx(path, toy.blocks)
            engine.ingest(
                path,
                tenant_id=tenant_id,
                course_id=toy.course_id,
                document_id=toy.document_id,
            )

        for question in QUESTIONS:
            scope = Scope(tenant_id=tenant_id, course_ids=frozenset({question.course_id}))
            results, diagnostics = engine.retrieve(scope, question.question, top_k)
            retrieved = [r.chunk.document_id for r in results]
            row = {
                "qid": question.qid,
                "language": question.language,
                "case_type": question.case_type,
                "retrieved": retrieved,
                "top_score": round(diagnostics.top_score, 4),
                "hit": hit_at_k(retrieved, list(question.gold_docs)) if question.gold_docs else None,
                "recall": recall_at_k(retrieved, list(question.gold_docs)) if question.gold_docs else None,
                "answer_f1": None,
                "abstained": None,
            }
            if llm is not None:
                answer = engine.answer(question.question, scope, top_k)
                row["abstained"] = answer.abstained
                if question.reference:
                    row["answer_f1"] = round(token_f1(answer.answer, question.reference), 4)
            rows.append(row)

    scored = [r for r in rows if r["hit"] is not None]
    summary = {
        "embedder": embedder_label,
        "llm": llm_label,
        "top_k": top_k,
        "questions": len(rows),
        "scored_questions": len(scored),
        "hit_at_k": round(sum(r["hit"] for r in scored) / len(scored), 4) if scored else None,
        "recall_at_k": round(sum(r["recall"] for r in scored) / len(scored), 4) if scored else None,
        "rows": rows,
    }
    return summary
