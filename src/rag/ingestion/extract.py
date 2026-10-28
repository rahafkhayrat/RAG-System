from __future__ import annotations

from pathlib import Path

from ..contracts import ExtractedDocument, ExtractedPage


def extract_document(
    path: str | Path,
    *,
    tenant_id: str,
    course_id: str,
    document_id: str,
    source_uri: str | None = None,
) -> ExtractedDocument:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(str(p))

    suffix = p.suffix.lower()
    try:
        if suffix == ".pdf":
            pages = _extract_pdf(p)
        elif suffix == ".docx":
            pages = _extract_docx(p)
        else:
            raise ValueError(f"unsupported file type: {suffix or '<none>'}")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"could not read {p.name}: invalid, damaged, or unsupported document") from exc

    pages = tuple(pg for pg in pages if pg.text.strip())
    if not pages:
        raise ValueError(f"no extractable text in {p.name}")

    return ExtractedDocument(
        tenant_id=tenant_id,
        course_id=course_id,
        document_id=document_id,
        filename=p.name,
        source_uri=source_uri or p.name,
        pages=pages,
    )


def _extract_pdf(path: Path) -> list[ExtractedPage]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[ExtractedPage] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append(ExtractedPage(page=number, text=text.strip(), section=None))
    return pages


def _extract_docx(path: Path) -> list[ExtractedPage]:
    import docx

    document = docx.Document(str(path))
    blocks: list[ExtractedPage] = []
    current_section: str | None = None
    current_lines: list[str] = []

    def flush() -> None:
        text = "\n".join(current_lines).strip()
        if text:
            blocks.append(ExtractedPage(page=1, text=text, section=current_section))

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style else ""
        if style.startswith("Heading"):
            flush()
            current_lines = []
            current_section = text
        else:
            current_lines.append(text)
    flush()
    return blocks
