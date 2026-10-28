import docx

from rag.contracts import Chunk, RetrievedChunk, Scope
from rag.embeddings.hashing import HashingEmbedder
from rag.pipeline import RagEngine
from rag.store.memory import InMemoryVectorStore


# --------------------------------------------------
# Test helpers
# --------------------------------------------------

class FakeLLM:
    """Fake LLM that avoids external API calls."""

    def __init__(self, response: str):
        self.response = response
        self.calls = []

    def complete(self, system_prompt: str, user_prompt: str) -> str:
        self.calls.append(
            {
                "system_prompt": system_prompt,
                "user_prompt": user_prompt,
            }
        )
        return self.response


class FakeSufficiencyChecker:
    """Fake checker for testing evidence sufficiency behavior."""

    def __init__(self, sufficient: bool, reason: str = ""):
        self.sufficient = sufficient
        self.reason = reason
        self.calls = []

    def evaluate(self, question, evidence):
        from rag.evaluation.sufficiency import SufficiencyResult

        self.calls.append(
            {
                "question": question,
                "evidence": evidence,
            }
        )

        return SufficiencyResult(
            sufficient=self.sufficient,
            reason=self.reason,
            evidence_count=len(evidence),
        )


def make_evidence(text: str) -> RetrievedChunk:
    chunk = Chunk(
        tenant_id="t1",
        course_id="math",
        document_id="doc1",
        chunk_id="doc1#0",
        text=text,
        source_uri="math.docx",
        page=1,
    )

    return RetrievedChunk(
        chunk=chunk,
        score=0.9,
    )


def _write_docx(tmp_path, name, heading, body):
    path = tmp_path / name

    document = docx.Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(body)
    document.save(str(path))

    return path


# --------------------------------------------------
# Original pipeline tests
# --------------------------------------------------

def test_ingest_and_scoped_retrieval(tmp_path):
    engine = RagEngine(
        store=InMemoryVectorStore(),
        embedder=HashingEmbedder(dim=128),
    )

    path_a = _write_docx(
        tmp_path,
        "a.docx",
        "Algebra",
        "linear equations a x plus b equals zero",
    )

    path_b = _write_docx(
        tmp_path,
        "b.docx",
        "History",
        "industrial revolution steam power",
    )

    assert engine.ingest(
        path_a,
        tenant_id="t1",
        course_id="math",
        document_id="a",
    ) >= 1

    assert engine.ingest(
        path_b,
        tenant_id="t1",
        course_id="hist",
        document_id="b",
    ) >= 1

    scope = Scope(
        tenant_id="t1",
        course_ids=frozenset({"math"}),
    )

    results, diagnostics = engine.retrieve(
        scope,
        "linear equations",
        top_k=5,
    )

    assert results
    assert all(
        result.chunk.document_id == "a"
        for result in results
    )
    assert diagnostics.returned == len(results)


def test_other_tenant_never_retrieved(tmp_path):
    engine = RagEngine(
        store=InMemoryVectorStore(),
        embedder=HashingEmbedder(dim=128),
    )

    path = _write_docx(
        tmp_path,
        "a.docx",
        "Algebra",
        "linear equations a x plus b equals zero",
    )

    engine.ingest(
        path,
        tenant_id="tenantA",
        course_id="math",
        document_id="a",
    )

    scope_b = Scope(
        tenant_id="tenantB",
        course_ids=frozenset({"math"}),
    )

    results, _ = engine.retrieve(
        scope_b,
        "linear equations",
        top_k=5,
    )

    assert results == []


def test_answer_abstains_without_evidence():
    engine = RagEngine(
        store=InMemoryVectorStore(),
        embedder=HashingEmbedder(dim=64),
    )

    scope = Scope(
        tenant_id="t1",
        course_ids=frozenset({"math"}),
    )

    result = engine.answer(
        "anything",
        scope,
    )

    assert result.abstained is True
    assert result.reason == "no_evidence"


# --------------------------------------------------
# Evidence sufficiency integration tests
# --------------------------------------------------

def test_insufficient_evidence_causes_abstention():
    llm = FakeLLM("This answer should not be generated.")

    checker = FakeSufficiencyChecker(
        sufficient=False,
        reason="insufficient_evidence",
    )

    engine = RagEngine(
        store=InMemoryVectorStore(),
        embedder=HashingEmbedder(dim=64),
        llm=llm,
        sufficiency_checker=checker,
    )

    scope = Scope(
        tenant_id="t1",
        course_ids=frozenset({"math"}),
    )

    evidence = make_evidence(
        "Linear algebra studies vectors."
    )

    engine.store.upsert(
        [evidence.chunk],
        engine.embedder.embed([evidence.chunk.text]),
    )

    result = engine.answer(
        "When is the mathematics exam?",
        scope,
    )

    assert result.abstained is True
    assert result.reason == "insufficient_evidence"
    assert result.sources


def test_generation_llm_is_not_called_when_evidence_is_insufficient():
    llm = FakeLLM("Generated answer")

    checker = FakeSufficiencyChecker(
        sufficient=False,
        reason="insufficient_evidence",
    )

    engine = RagEngine(
        store=InMemoryVectorStore(),
        embedder=HashingEmbedder(dim=64),
        llm=llm,
        sufficiency_checker=checker,
    )

    scope = Scope(
        tenant_id="t1",
        course_ids=frozenset({"math"}),
    )

    evidence = make_evidence(
        "Linear algebra studies vectors."
    )

    engine.store.upsert(
        [evidence.chunk],
        engine.embedder.embed([evidence.chunk.text]),
    )

    result = engine.answer(
        "When is the mathematics exam?",
        scope,
    )

    assert result.abstained is True
    assert result.reason == "insufficient_evidence"

    # The checker runs before and after the retry.
    assert len(checker.calls) == 2

    # The answer-generation LLM must not be called.
    assert llm.calls == []


def test_sufficient_evidence_allows_answer_generation():
    llm = FakeLLM(
        "Linear algebra studies vectors and linear equations."
    )

    checker = FakeSufficiencyChecker(
        sufficient=True,
        reason="Evidence directly answers the question.",
    )

    engine = RagEngine(
        store=InMemoryVectorStore(),
        embedder=HashingEmbedder(dim=64),
        llm=llm,
        sufficiency_checker=checker,
    )

    scope = Scope(
        tenant_id="t1",
        course_ids=frozenset({"math"}),
    )

    evidence = make_evidence(
        "Linear algebra studies vectors and linear equations."
    )

    engine.store.upsert(
        [evidence.chunk],
        engine.embedder.embed([evidence.chunk.text]),
    )

    result = engine.answer(
        "What does linear algebra study?",
        scope,
    )

    assert result.abstained is False
    assert result.answer == (
        "Linear algebra studies vectors and linear equations."
    )
    assert len(checker.calls) == 1
    assert len(llm.calls) == 1


def test_sufficiency_check_failure_causes_abstention():
    llm = FakeLLM("Generated answer")

    checker = FakeSufficiencyChecker(
        sufficient=False,
        reason="sufficiency_evaluation_failed",
    )

    engine = RagEngine(
        store=InMemoryVectorStore(),
        embedder=HashingEmbedder(dim=64),
        llm=llm,
        sufficiency_checker=checker,
    )

    scope = Scope(
        tenant_id="t1",
        course_ids=frozenset({"math"}),
    )

    evidence = make_evidence(
        "Some course content."
    )

    engine.store.upsert(
        [evidence.chunk],
        engine.embedder.embed([evidence.chunk.text]),
    )

    result = engine.answer(
        "Explain the topic.",
        scope,
    )

    assert result.abstained is True
    assert result.reason == "sufficiency_evaluation_failed"
    assert llm.calls == []