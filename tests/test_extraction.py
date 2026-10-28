import docx
import pytest

from rag.ingestion.extract import extract_document


def test_extract_docx_preserves_sections(tmp_path):
    path = tmp_path / "course.docx"
    document = docx.Document()
    document.add_heading("Linear Equations", level=1)
    document.add_paragraph("An equation of the form a x + b = 0.")
    document.add_heading("Quadratics", level=1)
    document.add_paragraph("a x squared + b x + c = 0.")
    document.save(str(path))

    extracted = extract_document(path, tenant_id="t1", course_id="c1", document_id="doc1")

    assert extracted.document_id == "doc1"
    assert extracted.course_id == "c1"
    sections = [page.section for page in extracted.pages]
    assert "Quadratics" in sections
    assert any("a x + b = 0" in page.text for page in extracted.pages)


def test_extract_rejects_unsupported_type(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError):
        extract_document(path, tenant_id="t1", course_id="c1", document_id="doc1")


def test_extract_rejects_empty_document(tmp_path):
    path = tmp_path / "empty.docx"
    docx.Document().save(str(path))
    with pytest.raises(ValueError):
        extract_document(path, tenant_id="t1", course_id="c1", document_id="doc1")


def test_extract_rejects_damaged_docx_with_clear_value_error(tmp_path):
    path = tmp_path / "damaged.docx"
    path.write_text("not a real docx", encoding="utf-8")
    with pytest.raises(ValueError, match="could not read damaged.docx"):
        extract_document(path, tenant_id="t1", course_id="c1", document_id="doc1")
