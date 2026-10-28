from types import SimpleNamespace
from io import BytesIO

import docx
import rag.api.app as api
from fastapi.testclient import TestClient
from rag.contracts import Scope
from rag.embeddings.hashing import HashingEmbedder
from rag.evaluation.sufficiency import SufficiencyResult
from rag.pipeline import RagEngine
from rag.store.memory import InMemoryVectorStore


class FakeEngine:
    def __init__(self, chunks=2):
        self.chunks = chunks
        self.llm = None
        self.ingested = []
        self.answered = []

    def ingest(self, path, **kwargs):
        self.ingested.append(kwargs)
        return self.chunks

    def answer(self, question, scope):
        self.answered.append((question, scope))
        return SimpleNamespace(
            question=question,
            answer="The course material says this is a test answer.",
            abstained=False,
            reason="",
            sources=[],
        )


def make_client(tmp_path, monkeypatch, engine=None):
    monkeypatch.setattr(
        api.Settings,
        "from_env",
        lambda: SimpleNamespace(data_dir=tmp_path, top_k=3),
    )
    app = api.create_app(engine or FakeEngine())
    return TestClient(app)


def test_health_endpoint(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)

    response = client.get("/api/health")

    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_upload_rejects_missing_course_id(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)

    response = client.post(
        "/api/documents",
        files={
            "file": ("notes.pdf", BytesIO(b"fake document"), "application/pdf")
        },
    )

    assert response.status_code == 400
    assert "course_id" in response.json()["error"]


def test_upload_rejects_requests_over_size_limit(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)

    response = client.post(
        "/api/documents",
        content=b"x",
        headers={"content-length": str(api.MAX_UPLOAD_BYTES + 1)},
    )

    assert response.status_code == 413
    assert response.json()["error"] == (
        "The total upload request exceeds the 20 MB limit."
    )


def test_upload_rejects_unsupported_extension(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)

    response = client.post(
        "/api/documents",
        data={"course_id": "math101"},
        files={
            "file": ("notes.txt", BytesIO(b"fake document"), "text/plain")
        },
    )

    assert response.status_code == 415


def test_upload_returns_indexing_result(tmp_path, monkeypatch):
    engine = FakeEngine(chunks=2)
    client = make_client(tmp_path, monkeypatch, engine)

    response = client.post(
        "/api/documents",
        data={
            "course_id": "math101",
        },
        files={
            "file": ("notes.pdf", BytesIO(b"fake PDF content"), "application/pdf")
        },
    )

    assert response.status_code == 201
    assert response.json()["filename"] == "notes.pdf"
    assert response.json()["course_id"] == "math101"
    assert response.json()["chunks"] == 2
    assert len(engine.ingested) == 1


def test_question_requires_valid_json(tmp_path, monkeypatch):
    client = make_client(tmp_path, monkeypatch)

    response = client.post("/api/questions", content="not json")

    assert response.status_code == 400


def test_question_returns_answer(tmp_path, monkeypatch):
    engine = FakeEngine()
    client = make_client(tmp_path, monkeypatch, engine)

    response = client.post(
        "/api/questions",
        json={
            "course_id": "math101",
            "question": "What is covered in the course?",
        },
    )

    assert response.status_code == 200
    assert response.json()["answer"] == (
        "The course material says this is a test answer."
    )
    assert response.json()["abstained"] is False
    assert len(engine.answered) == 1


class _MathAnswerLLM:
    def complete(self, _system, _user):
        return r"The condition is \( \det(A - \lambda_1 I) = 0 \). [1]"


class _SelectiveMathSufficiencyChecker:
    def evaluate(self, question, evidence):
        supported = "lambda_1" in question.casefold() and bool(evidence)
        return SufficiencyResult(
            sufficient=supported,
            reason="test evidence accepted" if supported else "insufficient_evidence",
            evidence_count=len(evidence),
        )


def test_real_upload_math_retrieval_and_course_scope(tmp_path, monkeypatch):
    engine = RagEngine(
        store=InMemoryVectorStore(),
        embedder=HashingEmbedder(dim=128),
        llm=_MathAnswerLLM(),
        sufficiency_checker=_SelectiveMathSufficiencyChecker(),
    )
    client = make_client(tmp_path, monkeypatch, engine)

    document = docx.Document()
    document.add_heading("Eigenvalues", level=1)
    document.add_paragraph(
        "The scalar λ₁ satisfies det(A − λ₁I) = 0. "
        "The polynomial x² = 4 has roots 2 and −2."
    )
    content = BytesIO()
    document.save(content)
    content.seek(0)

    upload = client.post(
        "/api/documents",
        data={"course_id": "math101"},
        files={
            "file": (
                "math-notes.docx",
                content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert upload.status_code == 201
    assert upload.json()["chunks"] == 1

    scope = Scope("local-user", frozenset({"math101"}))
    keyword_results = engine.store.search_keywords(
        scope,
        "What does lambda_1 satisfy?",
        top_k=1,
    )
    assert keyword_results
    assert keyword_results[0].chunk.text.startswith("The scalar λ₁")

    answer = client.post(
        "/api/questions",
        json={"course_id": "math101", "question": "What does lambda_1 satisfy?"},
    )

    assert answer.status_code == 200
    assert answer.json()["abstained"] is False
    assert answer.json()["sources"][0]["filename"] == "math-notes.docx"
    assert answer.json()["sources"][0]["page"] == 1

    unrelated = client.post(
        "/api/questions",
        json={"course_id": "math101", "question": "When is the final exam?"},
    )
    assert unrelated.json()["abstained"] is True
    assert unrelated.json()["reason"] == "insufficient_evidence"

    other_course = client.post(
        "/api/questions",
        json={"course_id": "other-course", "question": "What does lambda_1 satisfy?"},
    )
    assert other_course.json()["abstained"] is True
    assert other_course.json()["reason"] == "no_evidence"
    assert other_course.json()["sources"] == []
